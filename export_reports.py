
from pathlib import Path
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# --------------------------------------------------
# 1. Project paths
# --------------------------------------------------
PROJECT_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_DIR / "data" / "processed"
REPORT_DIR = PROJECT_DIR / "reports"

REPORT_DIR.mkdir(parents=True, exist_ok=True)

EMPLOYEE_FILE = DATA_DIR / "employees_clean.csv"
MONTHLY_FILE = DATA_DIR / "monthly_headcount.csv"
QUALITY_FILE = DATA_DIR / "data_quality_summary.csv"

EXCEL_FILE = REPORT_DIR / "workforce_report.xlsx"
DEPT_FILE = DATA_DIR / "monthly_department_summary.csv"

# --------------------------------------------------
# 2. Load existing outputs
# --------------------------------------------------
employees = pd.read_csv(EMPLOYEE_FILE)
monthly = pd.read_csv(MONTHLY_FILE)
quality = pd.read_csv(QUALITY_FILE)

employees["DateofHire"] = pd.to_datetime(
    employees["DateofHire"], errors="coerce"
)
employees["DateofTermination"] = pd.to_datetime(
    employees["DateofTermination"], errors="coerce"
)

monthly["Month"] = monthly["Month"].astype(str)

# --------------------------------------------------
# 3. Calculate department-level monthly metrics
# --------------------------------------------------
departments = sorted(employees["Department"].dropna().unique())
department_rows = []

for month_text in monthly["Month"]:
    month = pd.Period(month_text, freq="M")
    month_start = month.start_time.normalize()
    month_end = month.end_time.normalize()

    for department in departments:
        dept = employees[
            employees["Department"] == department
        ]

        # Opening balance: employed immediately before
        # the first day of this month.
        active_at_start = (
            (dept["DateofHire"] < month_start)
            & (
                dept["DateofTermination"].isna()
                | (dept["DateofTermination"] >= month_start)
            )
        )

        # Closing balance: employed at the end of the month.
        active_at_end = (
            (dept["DateofHire"] <= month_end)
            & (
                dept["DateofTermination"].isna()
                | (dept["DateofTermination"] > month_end)
            )
        )

        hires = (
            (dept["DateofHire"] >= month_start)
            & (dept["DateofHire"] <= month_end)
        ).sum()

        exits = (
            (dept["DateofTermination"] >= month_start)
            & (dept["DateofTermination"] <= month_end)
        ).sum()

        opening = int(active_at_start.sum())
        closing = int(active_at_end.sum())
        average = (opening + closing) / 2

        department_rows.append({
            "Month": str(month),
            "Department": department,
            "OpeningHeadcount": opening,
            "Hires": int(hires),
            "Exits": int(exits),
            "ClosingHeadcount": closing,
            "NetChange": int(hires - exits),
            "AttritionRate": exits / average if average else None,
            "ReconciliationPassed": (
                opening + hires - exits == closing
            ),
        })

department_monthly = pd.DataFrame(department_rows)

# Save the Power BI input file
department_monthly.to_csv(DEPT_FILE, index=False)

# --------------------------------------------------
# 4. Build the executive summary
# --------------------------------------------------
latest = monthly.iloc[-1]

summary = pd.DataFrame({
    "Metric": [
        "Reporting period",
        "Latest closing headcount",
        "Total hires in dataset",
        "Total exits in dataset",
        "Months validated",
        "Department-month checks passed",
        "Employee records",
        "Duplicate employee IDs",
        "Missing manager IDs",
    ],
    "Value": [
        monthly["Month"].iloc[-1],
        int(latest["ClosingHeadcount"]),
        int(monthly["Hires"].sum()),
        int(monthly["Exits"].sum()),
        f'{int(monthly["ReconciliationPassed"].sum())} / {len(monthly)}',
        f'{int(department_monthly["ReconciliationPassed"].sum())} / {len(department_monthly)}',
        len(employees),
        int(employees["EmpID"].duplicated().sum()),
        int(employees["ManagerID"].isna().sum()),
    ],
})

# Keep only fields required for the reporting workflow.
employee_columns = [
    "EmpID",
    "Department",
    "Position",
    "DateofHire",
    "DateofTermination",
    "EmploymentStatus",
    "ManagerID",
    "ManagerName",
]

employee_report = employees[
    [c for c in employee_columns if c in employees.columns]
].copy()

# --------------------------------------------------
# 5. Create the Excel workbook
# --------------------------------------------------
with pd.ExcelWriter(EXCEL_FILE, engine="openpyxl") as writer:
    summary.to_excel(
        writer, sheet_name="Executive Summary", index=False
    )

    monthly.to_excel(
        writer, sheet_name="Monthly Headcount", index=False
    )

    department_monthly.to_excel(
        writer, sheet_name="Department Analysis", index=False
    )

    quality.to_excel(
        writer, sheet_name="Data Quality", index=False
    )

    employee_report.to_excel(
        writer, sheet_name="Employee Data", index=False
    )

# --------------------------------------------------
# 6. Format the workbook
# --------------------------------------------------
from openpyxl import load_workbook

workbook = load_workbook(EXCEL_FILE)

for sheet in workbook.worksheets:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    # Format header row
    for cell in sheet[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="244062"
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # Set reasonable column widths
    for column_cells in sheet.columns:
        column_letter = get_column_letter(
            column_cells[0].column
        )

        max_length = max(
            len(str(cell.value or ""))
            for cell in column_cells
        )

        sheet.column_dimensions[column_letter].width = min(
            max(max_length + 2, 12), 32
        )

    # Format attrition as a percentage
    header_map = {
        cell.value: cell.column
        for cell in sheet[1]
    }

    if "AttritionRate" in header_map:
        col = header_map["AttritionRate"]

        for row in range(2, sheet.max_row + 1):
            sheet.cell(row=row, column=col).number_format = "0.00%"

    # Format dates
    for column_name in ["DateofHire", "DateofTermination"]:
        if column_name in header_map:
            col = header_map[column_name]

            for row in range(2, sheet.max_row + 1):
                sheet.cell(
                    row=row, column=col
                ).number_format = "yyyy-mm-dd"

workbook.save(EXCEL_FILE)

# --------------------------------------------------
# 7. Report completion
# --------------------------------------------------
print("\n--- REPORT GENERATION COMPLETE ---")
print(f"Excel report: {EXCEL_FILE}")
print(f"Department CSV: {DEPT_FILE}")
print(f"Employee records exported: {len(employee_report)}")
print(f"Department-month rows: {len(department_monthly)}")

if department_monthly["ReconciliationPassed"].all():
    print("All department-month headcount checks passed.")
else:
    print("WARNING: Review department reconciliation errors.")
